from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Tasia_Chernyakhovskaya_CV.docx"
FONT = "Arial"
INK = RGBColor(23, 23, 23)
MUTED = RGBColor(102, 102, 102)
LINK = RGBColor(80, 62, 220)


def set_run_font(run, size=None, bold=None, color=INK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color
    return run


def set_spacing(paragraph, before=0, after=0, line=1.05):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.widow_control = True


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relation = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run_element = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "503EDC")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    properties.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    properties.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    properties.append(size)
    run_element.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run_element.append(text_element)
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)


def add_section_heading(doc, text):
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    set_spacing(paragraph, before=9, after=4, line=1)
    set_run_font(paragraph.add_run(text), 14, True)
    return paragraph


def add_job(doc, company, role, dates, description, bullets):
    heading = doc.add_paragraph()
    heading.paragraph_format.keep_with_next = True
    set_spacing(heading, before=4.5, after=0.8, line=1.04)
    set_run_font(heading.add_run(company), 11, True)
    set_run_font(heading.add_run("  |  " + role), 11, True)

    meta = doc.add_paragraph()
    meta.paragraph_format.keep_with_next = True
    set_spacing(meta, after=1.5, line=1)
    set_run_font(meta.add_run(dates), 9.6, False, MUTED)

    if description:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        set_spacing(paragraph, after=2, line=1.04)
        set_run_font(paragraph.add_run(description), 9.7)

    for item in bullets:
        paragraph = doc.add_paragraph(style="CV Bullet")
        set_spacing(paragraph, after=1, line=1.04)
        set_run_font(paragraph.add_run("• "), 9.6)
        set_run_font(paragraph.add_run(item), 9.6)


def add_skill(doc, label, value):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, after=1.6, line=1.04)
    set_run_font(paragraph.add_run(label + ": "), 9.7, True)
    set_run_font(paragraph.add_run(value), 9.7)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(9.7)
    normal.font.color.rgb = INK

    title_style = doc.styles["Title"]
    title_style.font.name = FONT
    title_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.font.color.rgb = INK
    title_ppr = title_style.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    heading_style = doc.styles["Heading 1"]
    heading_style.font.name = FONT
    heading_style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    heading_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    heading_style.font.color.rgb = INK

    if "CV Bullet" not in doc.styles:
        bullet_style = doc.styles.add_style("CV Bullet", WD_STYLE_TYPE.PARAGRAPH)
    else:
        bullet_style = doc.styles["CV Bullet"]
    bullet_style.base_style = normal
    bullet_style.paragraph_format.left_indent = Cm(0.45)
    bullet_style.paragraph_format.first_line_indent = Cm(-0.32)

    title = doc.add_paragraph(style="Title")
    set_spacing(title, after=1, line=1)
    set_run_font(title.add_run("Таисия Черняховская"), 25, True)

    subtitle = doc.add_paragraph()
    set_spacing(subtitle, after=2.5, line=1)
    set_run_font(subtitle.add_run("Senior Marketing & Web Designer"), 12.5, True)

    contacts = doc.add_paragraph()
    set_spacing(contacts, after=5, line=1)
    set_run_font(contacts.add_run("Екатеринбург  |  GMT+5  |  удалённо  |  "), 9.6, False, MUTED)
    add_hyperlink(contacts, "Портфолио", "https://tallety.ru")
    set_run_font(contacts.add_run("  |  "), 9.4, False, MUTED)
    add_hyperlink(contacts, "Telegram", "https://t.me/tallety")
    set_run_font(contacts.add_run("  |  "), 9.4, False, MUTED)
    add_hyperlink(contacts, "heytallety@gmail.com", "mailto:heytallety@gmail.com")

    summary = doc.add_paragraph()
    set_spacing(summary, after=4, line=1.1)
    set_run_font(
        summary.add_run(
            "Senior Marketing & Web Designer с 4 годами опыта в EdTech, цифровых коммуникациях и брендинге. "
            "Разрабатываю key visual и масштабирую его на рекламные креативы, контент, презентации, "
            "лендинги и печатные носители. Работаю с 3D-графикой и развиваю 2D-эскизы в объекты "
            "для key visual и веб-композиций."
        ),
        10,
    )
    summary_two = doc.add_paragraph()
    set_spacing(summary_two, after=4.5, line=1.1)
    set_run_font(
        summary_two.add_run(
            "В проекте «Башня» при МГТУ им. Н.Э. Баумана отвечаю за визуальное направление, шаблоны, "
            "документацию, дизайн-ревью и развитие младшего дизайнера. Работала над партнёрскими "
            "проектами с Яндексом, Ozon, Авито, Сбером, Т-Банком, VK и ВТБ."
        ),
        10,
    )

    add_section_heading(doc, "Опыт")
    add_job(
        doc,
        "Башня, МГТУ им. Н.Э. Баумана",
        "Ведущий коммуникационный дизайнер",
        "апрель 2024 - настоящее время",
        "EdTech-проект с курсами, конференциями и хакатонами совместно с российскими IT-компаниями.",
        [
            "Разработала 20+ визуальных концепций для образовательных продуктов, запусков и партнёрских мероприятий.",
            "Создала визуальную систему бренда: типографика, цвета, графические элементы и правила масштабирования.",
            "Собрала 15+ шаблонов для соцсетей, презентаций, рассылок и печатных материалов.",
            "Масштабировала key visual на digital-креативы, контент, лендинги и офлайн-носители.",
            "Работала с 3D-графикой: развивала 2D-эскизы в объекты для key visual и веб-композиций.",
            "Оформила 10+ партнёрских проектов с Яндексом, Ozon, Авито, Сбером, Т-Банком, VK и ВТБ.",
            "Выстроила дизайн-ревью, документацию и онбординг; курирую младшего дизайнера.",
        ],
    )
    add_job(
        doc,
        "Doggo & PuppyClub",
        "Коммуникационный дизайнер / иллюстратор",
        "сентябрь 2023 - апрель 2024",
        None,
        [
            "Разработала фирменных персонажей и визуальный язык сети груминг-салонов.",
            "Масштабировала систему на соцсети, рекламные креативы, стикерпаки, упаковку и полиграфию.",
            "Готовила цифровые и печатные макеты к производству.",
        ],
    )
    add_job(
        doc,
        "Фриланс и проектная работа",
        "Коммуникационный дизайнер / иллюстратор",
        "сентябрь 2022 - сентябрь 2023",
        None,
        [
            "Создавала визуальные концепции, сайты, рекламные и контентные материалы для брендов и студий.",
            "Самостоятельно вела проекты от брифа до подготовки файлов к публикации и печати.",
        ],
    )

    add_section_heading(doc, "Навыки")
    add_skill(doc, "Marketing design", "key visual, рекламные кампании, адаптации, digital-креативы, контентные системы")
    add_skill(doc, "Web и UI", "лендинги, прототипирование, адаптивные макеты, Framer, Tilda")
    add_skill(doc, "Figma", "шаблоны, визуальные библиотеки, прототипы и подготовка макетов")
    add_skill(doc, "Graphics", "Adobe Photoshop, Illustrator, Procreate, ретушь, коллаж, цветокоррекция, 3D-графика")
    add_skill(doc, "AI workflow", "Recraft, Krea, prompting, генерация серий, ручная доработка и визуальный QA")
    add_skill(doc, "Team", "арт-дирекшн, дизайн-ревью, документация, менторинг")
    add_skill(doc, "English", "B2")

    add_section_heading(doc, "Образование")
    education = doc.add_paragraph()
    set_spacing(education, after=0, line=1.05)
    set_run_font(education.add_run("СПГХПА им. А.Л. Штиглица"), 9.6, True)
    set_run_font(education.add_run("  |  54.03.01 Дизайн, профиль «Дизайн костюма»  |  2022-2026"), 9.4)

    doc.core_properties.title = "Таисия Черняховская Senior Marketing and Web Designer CV"
    doc.core_properties.subject = "CV"
    doc.core_properties.author = "Таисия Черняховская"
    doc.core_properties.keywords = "Senior Marketing Designer, Web Designer, EdTech, Figma, 3D"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
